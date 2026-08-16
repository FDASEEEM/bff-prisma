# -*- coding: utf-8 -*-
"""Genera el Informe de Pruebas Unitarias (.docx) del PRISMA BFF con datos REALES
de la ejecución de Jest (--coverage). Incluye gráficos (matplotlib) embebidos."""
import os
from datetime import date

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

BLUE = "#307FE2"
YELLOW = "#FFB800"
DARK = "#1A1A1A"
GREEN = "#43B02A"

HERE = os.path.dirname(os.path.abspath(__file__))
IMG = os.path.join(HERE, "_report_assets")
os.makedirs(IMG, exist_ok=True)
DESKTOP = os.path.join(os.path.expanduser("~"), "Desktop")
OUT = os.path.join(DESKTOP, "Informe_Pruebas_Unitarias_PRISMA_BFF.docx")

# ── Datos reales (jest --coverage) ───────────────────────────────────────────
SUMMARY = {"suites": "4 / 4", "tests": "35 / 35", "snapshots": "0", "time": "7.412 s"}
GLOBAL = [
    ("Sentencias (Statements)", 76, 458, 16.59),
    ("Ramas (Branches)", 12, 13, 92.30),
    ("Funciones (Functions)", 26, 178, 14.60),
    ("Líneas (Lines)", 68, 408, 16.66),
]
# módulo, %stmts, %branch, %funcs, %lines
MODULES = [
    ("dashboard.service.ts", 100.0, 100.0, 100.0, 100.0),
    ("microservice.client.ts", 93.54, 100.0, 77.77, 93.10),
    ("colegios.service.ts", 92.85, 100.0, 88.88, 91.66),
    ("auth.service.ts", 85.71, 80.0, 88.88, 84.21),
]
SUITES = [
    ("MicroserviceClient", 12),
    ("AuthService", 10),
    ("ColegiosService", 8),
    ("DashboardService", 5),
]
EXAMPLES = [
    ("MicroserviceClient (infraestructura)", [
        ("Construye la URL correcta por servicio (users → http://localhost:3001)", "PASS"),
        ("Propaga query params y el header Authorization al microservicio", "PASS"),
        ("GET / POST / PATCH / DELETE delegan en el HttpService con el método correcto", "PASS"),
        ("Lanza BadGatewayException ante error de red", "PASS"),
        ("Propaga respuestas 4xx sin lanzar excepción (passthrough)", "PASS"),
        ("Normaliza el path: antepone '/' si falta y no lo duplica", "PASS"),
    ]),
    ("AuthService (proxy a ms-users)", [
        ("login() llama a client.post('users', '/api/auth/login', {email, password})", "PASS"),
        ("register() / refresh() reenvían el payload correcto", "PASS"),
        ("logout() / me() / updateMe() adjuntan el authToken", "PASS"),
        ("extractToken() extrae el Bearer correctamente", "PASS"),
        ("extractToken() lanza UnauthorizedException si falta header o el esquema es inválido", "PASS"),
    ]),
    ("ColegiosService (proxy a ms-users)", [
        ("findAll / findOne / create / update / deactivate enrutan al endpoint y método correctos", "PASS"),
        ("getStats() y getProfessors() propagan id, query y authToken", "PASS"),
    ]),
    ("DashboardService (agregación cross-service)", [
        ("getUserDashboard() agrega datos de usuario + jobs en una sola respuesta", "PASS"),
        ("getColegioDashboard() combina colegio + stats + professors + consumo", "PASS"),
        ("Lanza BadGatewayException si algún servicio downstream falla", "PASS"),
    ]),
]

# ── Gráficos ────────────────────────────────────────────────────────────────
def chart_modules():
    names = [m[0] for m in MODULES][::-1]
    vals = [m[4] for m in MODULES][::-1]
    fig, ax = plt.subplots(figsize=(7.2, 3.0), dpi=150)
    bars = ax.barh(names, vals, color=BLUE)
    ax.set_xlim(0, 100)
    ax.set_xlabel("% Líneas cubiertas")
    ax.set_title("Cobertura de líneas por módulo con pruebas unitarias", fontsize=11, fontweight="bold")
    ax.axvline(70, color=YELLOW, linestyle="--", linewidth=1.5, label="Umbral sugerido 70%")
    for b, v in zip(bars, vals):
        ax.text(v - 3, b.get_y() + b.get_height() / 2, f"{v:.1f}%", va="center", ha="right", color="white", fontweight="bold", fontsize=9)
    ax.legend(loc="lower right", fontsize=8)
    fig.tight_layout()
    p = os.path.join(IMG, "modules.png")
    fig.savefig(p); plt.close(fig)
    return p

def chart_global():
    labels = ["Sentencias", "Ramas", "Funciones", "Líneas"]
    vals = [g[3] for g in GLOBAL]
    colors = [BLUE, GREEN, BLUE, BLUE]
    fig, ax = plt.subplots(figsize=(7.2, 3.0), dpi=150)
    bars = ax.bar(labels, vals, color=colors)
    ax.set_ylim(0, 100)
    ax.set_ylabel("% cubierto (global)")
    ax.set_title("Métricas globales de cobertura del proyecto", fontsize=11, fontweight="bold")
    for b, v in zip(bars, vals):
        ax.text(b.get_x() + b.get_width() / 2, v + 2, f"{v:.1f}%", ha="center", fontweight="bold", fontsize=9)
    fig.tight_layout()
    p = os.path.join(IMG, "global.png")
    fig.savefig(p); plt.close(fig)
    return p

def chart_suites():
    names = [s[0] for s in SUITES][::-1]
    vals = [s[1] for s in SUITES][::-1]
    fig, ax = plt.subplots(figsize=(7.2, 2.6), dpi=150)
    bars = ax.barh(names, vals, color=GREEN)
    ax.set_xlabel("N° de pruebas (todas PASS)")
    ax.set_title("Pruebas unitarias por suite (35 en total)", fontsize=11, fontweight="bold")
    for b, v in zip(bars, vals):
        ax.text(v + 0.15, b.get_y() + b.get_height() / 2, str(v), va="center", fontweight="bold", fontsize=9)
    ax.set_xlim(0, 14)
    fig.tight_layout()
    p = os.path.join(IMG, "suites.png")
    fig.savefig(p); plt.close(fig)
    return p

img_modules, img_global, img_suites = chart_modules(), chart_global(), chart_suites()

# ── Documento ────────────────────────────────────────────────────────────────
doc = Document()
doc.styles["Normal"].font.name = "Calibri"
doc.styles["Normal"].font.size = Pt(11)


def h(text, level=1):
    p = doc.add_heading(text, level=level)
    return p


def para(text, italic=False, size=None, color=None):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.italic = italic
    if size:
        r.font.size = Pt(size)
    if color:
        r.font.color.rgb = RGBColor.from_string(color.lstrip("#"))
    return p


def kv_table(rows, headers):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Light Grid Accent 1"
    for i, hd in enumerate(headers):
        c = t.rows[0].cells[i]
        c.paragraphs[0].add_run(hd).bold = True
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = str(val)
    return t


# Portada
title = doc.add_heading("Informe de Pruebas Unitarias", level=0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = sub.add_run("PRISMA BFF (Backend-for-Frontend) · prisma-bff")
r.bold = True
r.font.size = Pt(14)
r.font.color.rgb = RGBColor.from_string(BLUE.lstrip("#"))
meta = doc.add_paragraph()
meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
meta.add_run(f"Stack: NestJS 10 · TypeScript · Jest 29 + ts-jest    |    Fecha: {date.today().strftime('%d-%m-%Y')}").italic = True
doc.add_paragraph()

# 1. Introducción
h("1. Introducción y alcance", 1)
para("Este informe documenta las pruebas unitarias del microservicio PRISMA BFF "
     "(Backend-for-Frontend), la capa que agrega y orquesta las llamadas del frontend "
     "hacia los microservicios del sistema (ms-users, adminpanel, perfil-alumno, ms-docs). "
     "Las pruebas se centran en la lógica de negocio crítica: el cliente HTTP de "
     "microservicios, la capa de servicios (auth, colegios) y la agregación cross-service "
     "(dashboard). Todas las dependencias externas (HttpService, otros servicios) se "
     "aíslan con mocks, garantizando pruebas deterministas y sin red.")

# 2. Herramientas
h("2. Herramientas de testing", 1)
kv_table([
    ["Test runner", "Jest 29"],
    ["Transpilación TS", "ts-jest"],
    ["Utilidades NestJS", "@nestjs/testing (TestingModule, inyección de mocks)"],
    ["Cobertura de código", "Istanbul (integrado en Jest, flag --coverage)"],
    ["Reportes", "text, json-summary, lcov"],
    ["Comando", "npm test -- --coverage"],
], ["Herramienta", "Detalle"])

# 3. Resumen de ejecución
h("3. Resumen de ejecución", 1)
kv_table([
    ["Suites de prueba", SUMMARY["suites"] + " (100% passing)"],
    ["Pruebas (tests)", SUMMARY["tests"] + " (100% passing)"],
    ["Snapshots", SUMMARY["snapshots"]],
    ["Tiempo de ejecución", SUMMARY["time"]],
    ["Resultado", "✓ Todas las pruebas aprobadas"],
], ["Métrica", "Resultado"])
doc.add_paragraph()
doc.add_picture(img_suites, width=Inches(6.3))

# 4. Cobertura
h("4. Cobertura de código", 1)
para("Métricas globales generadas por Istanbul/Jest sobre todo el código fuente del BFF:")
kv_table([[g[0], f"{g[1]} / {g[2]}", f"{g[3]:.2f}%"] for g in GLOBAL],
         ["Métrica", "Cubierto / Total", "%"])
doc.add_paragraph()
doc.add_picture(img_global, width=Inches(6.3))
doc.add_paragraph()
para("La cobertura de ramas (branches) es muy alta (92.3%): la lógica condicional probada "
     "—manejo de errores, validación de tokens, normalización de URLs— está bien cubierta. "
     "La cobertura global de líneas/funciones es menor porque el alcance de las pruebas "
     "unitarias es la capa de servicios y el cliente de microservicios; los controllers "
     "(proxies finos) y los módulos de cableado de NestJS no se prueban a nivel unitario.")
doc.add_paragraph()
para("Cobertura de los módulos efectivamente bajo prueba unitaria:", italic=True)
kv_table([[m[0], f"{m[1]:.2f}%", f"{m[2]:.0f}%", f"{m[3]:.2f}%", f"{m[4]:.2f}%"] for m in MODULES],
         ["Módulo", "% Sentencias", "% Ramas", "% Funciones", "% Líneas"])
doc.add_paragraph()
doc.add_picture(img_modules, width=Inches(6.3))

# 5. Ejemplos
h("5. Ejemplos de pruebas realizadas y sus resultados", 1)
for suite, cases in EXAMPLES:
    h(suite, 2)
    for desc, res in cases:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(f"{desc}  ")
        rr = p.add_run(f"[{res}]")
        rr.bold = True
        rr.font.color.rgb = RGBColor.from_string(GREEN.lstrip("#"))

# 6. Conclusiones
h("6. Conclusiones y recomendaciones", 1)
for t in [
    "La lógica de negocio crítica del BFF está bien cubierta: el cliente de microservicios "
    "(93%), la agregación del dashboard (100%), el proxy de autenticación (84%) y la gestión "
    "de colegios (92%) superan ampliamente el umbral sugerido del 70% en líneas.",
    "Las 35 pruebas se ejecutan en ~7.4 s, son deterministas (sin red, todo mockeado) y "
    "pasan al 100%, lo que las hace aptas para integrarse como etapa bloqueante del pipeline CI/CD.",
    "Recomendación: incorporar pruebas a los controllers y a los servicios de jobs, students, "
    "professors y admin (hoy sin spec) para elevar la cobertura global de líneas por encima del 70%.",
    "Recomendación: agregar pruebas end-to-end (supertest, ya disponible como dependencia) para "
    "validar el ruteo HTTP completo del BFF.",
]:
    doc.add_paragraph(t, style="List Bullet")

doc.save(OUT)
print("OK ->", OUT)
